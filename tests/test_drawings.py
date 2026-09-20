"""Drawing sheets (plan / profile / cross-sections), SVG and DXF renderers, Excel workbook."""
import io
import xml.etree.ElementTree as ET

import numpy as np
import pytest

from plm.design.drawing.dxf import model_dxf, sheets_to_dxf
from plm.design.drawing.excel import design_workbook
from plm.design.drawing.frame import DrawingInputs, resolve_settings
from plm.design.drawing.model import Text, clip_polyline, readable
from plm.design.drawing.plan import plan_sheets, plan_windows
from plm.design.drawing.profile import profile_sheets, profile_windows
from plm.design.drawing.sections import section_sheets, split_regions
from plm.design.drawing.svg import sheet_to_svg
from plm.design.horizontal import curve_table
from plm.design.standards import load_standard
from plm.design.vertical import VerticalAlignment
from plm.engine.alignment import IP, HorizontalAlignment
from plm.engine.contour import ContourLine


def _alignment(length_x=800.0):
    return HorizontalAlignment([IP(0, 0, 0), IP(length_x * 0.45, 120, 150, "1", 40), IP(length_x, 0, 0)], 0.0)


def _section(ch, x, cut=True):
    """Two-lane design over a sloping ground: cut on the right, fill on the left when `cut`."""
    slope = 0.15 if cut else -0.15
    design = [(-9.0, 100 + 9 * slope * 0.6), (-4.0, 100.4), (-3.75, 100.1), (0.0, 100.2), (3.75, 100.1), (4.0, 100.4), (4.5, 99.9), (5.0, 99.9), (8.5, 100 + 8.5 * slope)]
    ground = [(o, 100 + o * slope) for o in np.linspace(-20, 20, 41)]
    return {"chainage": ch, "template_id": "default", "x": x, "y": 0.0, "direction": 0.0, "design_z": 100.2, "ground_z": 100.0, "slopes": [-2.5, -2.5],
            "design": design, "ground": ground, "cut_area": 3.0, "fill_area": 1.0, "flags": [],
            "left": {"kind": "fill" if cut else "cut", "hinge_offset": -4.0, "hinge_z": 100.4, "ground_z_at_hinge": 99.4, "height": 1.0, "catch_offset": -9.0, "benches": 0},
            "right": {"kind": "cut" if cut else "fill", "hinge_offset": 5.0, "hinge_z": 99.9, "ground_z_at_hinge": 100.75, "height": -0.85, "catch_offset": 8.5, "benches": 0}}


def _inputs(with_vertical=True, with_sections=True, n_sections=12, grade=0.04):
    al = _alignment()
    std = load_standard("nrs-2070")
    ctx = {"road_class": "feeder", "terrain": "rolling", "surface": "bituminous", "material": "soil", "design_speed": 40.0}
    L = al.length
    ground = [(c, 100 + grade * c + 2 * np.sin(c / 40.0)) for c in np.arange(0, L + 1, 5.0)]
    va = VerticalAlignment.make([{"chainage": 0, "elevation": 101, "length": 0}, {"chainage": L / 2, "elevation": 101 + grade * L / 2 + 1, "length": 60},
                                 {"chainage": L, "elevation": 101 + grade * L, "length": 0}]) if with_vertical else None
    secs = [_section(c, *al.point_at(c)[:1], cut=(i % 3 != 0)) for i, c in enumerate(np.linspace(0, L, n_sections))] if with_sections else []
    vols = [{"from": secs[i]["chainage"], "to": secs[i + 1]["chainage"], "length": secs[i + 1]["chainage"] - secs[i]["chainage"], "cut": 30.0, "fill": 10.0, "method": "end_area"} for i in range(len(secs) - 1)]
    mh = [{"chainage": secs[0]["chainage"], "cumulative": 0.0}] + [{"chainage": v["to"], "cumulative": 20.0 * (i + 1)} for i, v in enumerate(vols)] if secs else []
    contours = [ContourLine(level=float(z), is_major=(z % 5 == 0), coords=np.array([[x, 60 + (z - 100) * 10 + 10 * np.sin(x / 50)] for x in np.arange(-50, 900, 10.0)]), closed=False)
                for z in range(96, 112)]
    structures = [{"id": 1, "kind": "culvert", "side": None, "from": 300, "to": 300, "params": {"span": 0.9, "type": "pipe", "cells": 1}, "source": "manual", "note": ""},
                  {"id": 2, "kind": "retaining_wall", "side": "left", "from": 100, "to": 180, "params": {"height": 3.0, "type": "gabion"}, "source": "suggested", "note": ""},
                  {"id": 3, "kind": "side_drain", "side": "right", "from": 100, "to": 400, "params": {"width": 0.6, "depth": 0.5, "type": "masonry_trapezoid"}, "source": "suggested", "note": ""}]
    return DrawingInputs(project_name="Test project", design_name="Test road", al=al, curve_rows=curve_table(al, std, ctx), va=va, ground=ground, sections=secs,
                         volumes=vols, mass_haul=mh, totals={"cut": 330.0, "fill": 110.0, "net": 220.0}, structures=structures, contours=contours, supere=None,
                         standard=std.summary(), context=ctx, checks=[], standard_params=[], templates={}, crs_name="local", corridor_params={"interval": 20, "cut_factor": 1.0, "fill_factor": 1.0})


# ---------------------------------------------------------------------- helpers
def test_split_regions_crossing_ground():
    design = [(-5.0, 0.0), (5.0, 0.0)]
    ground = [(-5.0, -1.0), (5.0, 1.0)]
    regions = split_regions(design, ground)
    kinds = [k for k, _ in regions]
    assert kinds == ["fill", "cut"]

    def area(ring):
        x, y = zip(*ring)
        return 0.5 * abs(sum(x[i] * y[(i + 1) % len(x)] - x[(i + 1) % len(x)] * y[i] for i in range(len(x))))

    assert all(abs(area(r) - 2.5) < 1e-9 for _, r in regions)


def test_clip_and_readable():
    pieces = clip_polyline(np.array([[-5, 5], [5, 5], [15, 5]]), 0, 0, 10, 10)
    assert len(pieces) == 1 and np.allclose(pieces[0][0], [0, 5]) and np.allclose(pieces[0][-1], [10, 5])
    assert clip_polyline(np.array([[20, 20], [30, 30]]), 0, 0, 10, 10) == []
    assert readable(180.0, "left") == (0.0, "right") and readable(45.0, "left") == (45.0, "left")


def test_settings_resolution():
    st = resolve_settings({"paper": "a1", "plan_scale": 500, "fields": {"drawn_by": "PS"}, "kinds": ["plan"]}, project_name="P", design_name="D", organisation="Org")
    assert st.paper == "A1" and st.size == (841.0, 594.0) and st.plan_scale == 500 and st.kinds == ("plan",)
    assert st.fields["drawn_by"] == "PS" and st.fields["project"] == "P" and st.fields["organisation"] == "Org" and st.fields["road"] == "D"
    assert resolve_settings({"paper": "letter"}).paper == "A3"
    x0, y0, x1, y1 = st.area
    assert y0 > st.frame[1] and x1 - x0 > 700


# ---------------------------------------------------------------------- plan
def test_plan_windows_tile_the_alignment():
    inp = _inputs(with_sections=False)
    st = resolve_settings({"paper": "A3", "plan_scale": 1000})
    frames = plan_windows(inp.al, st)
    assert len(frames) >= 2
    assert frames[0].s0 == inp.al.start_chainage and abs(frames[-1].s1 - inp.al.end_chainage) < 1e-6
    for a, b in zip(frames, frames[1:]):
        assert abs(a.s1 - b.s0) < 1e-6
    # every window's alignment part fits inside the drawing area
    for fr in frames:
        xy, ch = inp.al.densify(5.0)
        m = (ch >= fr.s0) & (ch <= fr.s1)
        sxy = fr.to_sheet(np.asarray(xy)[m])
        x0, y0, x1, y1 = st.area
        assert sxy[:, 0].min() >= x0 - 1e-6 and sxy[:, 0].max() <= x1 + 1e-6 and sxy[:, 1].min() >= y0 - 1e-6 and sxy[:, 1].max() <= y1 + 1e-6
    # a big paper takes the whole road on one sheet
    assert len(plan_windows(inp.al, resolve_settings({"paper": "A0", "plan_scale": 1000}))) == 1
    # 1:50 on A4: the labels alone are wider than the paper
    with pytest.raises(ValueError):
        plan_windows(inp.al, resolve_settings({"paper": "A4", "plan_scale": 50}))


def test_plan_sheets_content():
    inp = _inputs()
    st = resolve_settings({"paper": "A3", "plan_scale": 1000, "fields": {"drawing_prefix": "T"}})
    sheets = plan_sheets(inp, st)
    assert len(sheets) >= 2 and sheets[0].number == "T-P-01" and sheets[0].kind == "plan"
    texts = [e.text for s in sheets for e in s.entities if isinstance(e, Text)]
    assert any("MATCH LINE" in t and "T-P-02" in t for t in texts)
    assert any(t.startswith("0+100") for t in texts)               # chainage labels
    assert any("CULVERT" in t for t in texts) and any("RETAINING WALL" in t and "Gabion" in t for t in texts)
    assert any("SIDE DRAIN" in t for t in texts)
    assert {"WALL", "DRAIN", "CULVERT"} <= {getattr(e, "layer", None) for s in sheets for e in s.entities}
    assert any(t in ("TS 0+" + x for x in [""]) or t.startswith("TS 0+") for t in texts)   # transition key points
    assert any("HORIZONTAL CURVE DATA" == t for t in texts)
    layers = {getattr(e, "layer", None) for s in sheets for e in s.entities}
    assert {"CONTOUR", "INDEX_CONTOUR", "H_ALIGN", "DAYLIGHT_CUT", "FORMATION", "FRAME", "TITLE"} <= layers
    # without terrain no contours, everything else the same
    lite = plan_sheets(inp, st, with_terrain=False)
    assert len(lite) == len(sheets) and not any(getattr(e, "layer", None) == "CONTOUR" for e in lite[0].entities)


# ---------------------------------------------------------------------- profile
def test_profile_windows_and_datum_segments():
    inp = _inputs(grade=0.12)
    st = resolve_settings({"paper": "A3", "profile_h_scale": 1000, "profile_v_scale": 100})
    windows = profile_windows(inp, st)
    # sheets cover the full drawing width; a steep road gets several datum segments per sheet
    assert len(windows) == 3 and windows[0][0] == 0 and abs(windows[-1][1] - inp.al.end_chainage) < 1e-6
    assert all(abs((b - a) - (windows[0][1] - windows[0][0])) < 1e-6 for a, b in windows[:-1])
    sheets = profile_sheets(inp, st)
    assert len(sheets) == len(windows)
    for s, (a, b) in zip(sheets, windows):
        assert s.info["start"] == a and s.info["end"] == b
        zmin = min(z for c, z in inp.ground if a <= c <= b)
        assert s.info["datum"] <= zmin and len(s.info["datums"]) >= 2
    flatter = profile_sheets(inp, resolve_settings({"paper": "A3", "profile_h_scale": 1000, "profile_v_scale": 200}))
    assert len(flatter[0].info["datums"]) < len(sheets[0].info["datums"])
    texts = [e.text for e in sheets[0].entities if isinstance(e, Text)]
    assert sum(t.startswith("DATUM RL") for t in texts) == len(sheets[0].info["datums"])
    assert any(t.startswith("DATUM CHANGE") for t in texts) and any("PVI" in t for t in texts) and "VERTICAL CURVE DATA" in texts
    assert any("%" in t for t in texts)
    # gentle road on a big sheet: a single datum, no break line
    gentle = profile_sheets(_inputs(grade=0.01), resolve_settings({"paper": "A1"}))
    assert len(gentle[0].info["datums"]) == 1 and not any(e.text.startswith("DATUM CHANGE") for e in gentle[0].entities if isinstance(e, Text))


def test_profile_without_vertical_still_draws_ground():
    inp = _inputs(with_vertical=False, with_sections=False)
    sheets = profile_sheets(inp, resolve_settings({"paper": "A1"}))
    assert sheets and any(getattr(e, "layer", None) == "GROUND" for e in sheets[0].entities)
    assert not any(getattr(e, "layer", None) == "DESIGN" for e in sheets[0].entities)


# ---------------------------------------------------------------------- sections
def test_section_sheets_grid_and_fallback_scale():
    inp = _inputs(n_sections=14)
    sheets = section_sheets(inp, resolve_settings({"paper": "A3", "section_scale": 200}))
    assert sum(s.info["sections"] for s in sheets) == 14 and sheets[0].kind == "sections"
    texts = [e.text for e in sheets[0].entities if isinstance(e, Text)]
    assert any(t.startswith("CH 0+000") for t in texts) and any(t.startswith("DATUM") for t in texts) and any(t.startswith("CUT ") for t in texts)
    layers = {getattr(e, "layer", None) for e in sheets[0].entities}
    assert {"CUT_HATCH", "FILL_HATCH", "GROUND", "DESIGN"} <= layers
    # the wall body and the drain channel are drawn on the sections they cover
    wall_sheets = [sh for sh in sheets if any(getattr(e, "layer", None) == "WALL" for e in sh.entities)]
    assert wall_sheets and any("Retaining wall" in e.text for sh in wall_sheets for e in sh.entities if isinstance(e, Text))
    assert any(getattr(e, "layer", None) == "DRAIN" for sh in sheets for e in sh.entities)
    # tiny paper at a large scale: the section is drawn smaller and says so instead of failing
    small = section_sheets(inp, resolve_settings({"paper": "A4", "section_scale": 50}))
    assert small and any("to fit" in e.text for s in small for e in s.entities if isinstance(e, Text))
    assert section_sheets(_inputs(with_sections=False), resolve_settings({})) == []


# ---------------------------------------------------------------------- renderers
def test_svg_is_well_formed_and_mirrors_sheet():
    inp = _inputs()
    sheet = plan_sheets(inp, resolve_settings({"paper": "A3"}), with_terrain=False)[0]
    svg = sheet_to_svg(sheet)
    root = ET.fromstring(svg)
    assert root.tag.endswith("svg") and root.get("viewBox") == "0 0 420 297" and root.get("data-kind") == "plan"
    ns = {"s": "http://www.w3.org/2000/svg"}
    n_text = len(root.findall(".//s:text", ns))
    assert n_text == sum(isinstance(e, Text) for e in sheet.entities) and n_text > 20
    assert "HORIZONTAL CURVE DATA" in svg and "pattern" in sheet_to_svg(section_sheets(inp, resolve_settings({}))[0])


def test_dxf_sheets_and_model(tmp_path):
    import ezdxf

    inp = _inputs(n_sections=6)
    st = resolve_settings({"paper": "A3"})
    sheets = plan_sheets(inp, st, with_terrain=False) + profile_sheets(inp, st)[:1] + section_sheets(inp, st)[:1]
    out = sheets_to_dxf(sheets, tmp_path / "sheets.dxf")
    doc = ezdxf.readfile(str(out))
    assert doc.header["$INSUNITS"] == 4
    names = doc.layouts.names()
    assert "P-01" in names and "L-01" in names and "X-01" in names
    assert {"H_ALIGN", "GROUND", "DESIGN", "CUT_HATCH", "TITLE_TEXT"} <= {ly.dxf.name for ly in doc.layers}
    msp = doc.modelspace()
    assert len(msp.query("LWPOLYLINE")) > 50 and len(msp.query("TEXT")) > 50 and len(msp.query("HATCH")) >= 2
    vp = doc.layouts.get("P-01").query("VIEWPORT")
    assert len(vp) >= 1

    m = model_dxf(inp, tmp_path / "model.dxf")
    mdoc = ezdxf.readfile(str(m))
    assert mdoc.header["$INSUNITS"] == 6
    cl = mdoc.modelspace().query('LWPOLYLINE[layer=="H_ALIGN"]')
    assert len(cl) == 1 and any(abs(b) > 0 for *_, b in cl[0].get_points("xyseb"))    # arcs as bulges
    assert len(mdoc.modelspace().query('POLYLINE[layer=="DESIGN_CL_3D"]')) == 1
    assert len(mdoc.modelspace().query('POLYLINE[layer=="DESIGN_SECTIONS_3D"]')) == 6
    assert len(mdoc.modelspace().query('TEXT[layer=="STRUCT_CULVERT"]')) == 1
    assert len(mdoc.modelspace().query('LWPOLYLINE[layer=="STRUCT_RETAINING_WALL"]')) == 1
    assert len(mdoc.modelspace().query('LWPOLYLINE[layer=="STRUCT_SIDE_DRAIN"]')) == 1


def test_excel_workbook():
    from openpyxl import load_workbook

    inp = _inputs(n_sections=5)
    data = design_workbook(inp, design={"name": "Test road", "settings": {}}, project={"name": "Test project", "crs": "local"})
    wb = load_workbook(io.BytesIO(data))
    assert {"Summary", "Horizontal", "Vertical", "Levels", "Sections", "Volumes", "Mass haul", "Section points", "Structures", "Checks", "Standard"} <= set(wb.sheetnames)
    ws = wb["Volumes"]
    assert ws["B1"].value == 1.0 and str(ws["J5"].value).startswith("=H5*$B$1") and str(ws["M6"].value).startswith("=M5+L6")
    assert wb["Horizontal"].max_row == 4 and wb["Structures"].max_row == 4
    assert "Structure quantities" in wb.sheetnames
    q = wb["Structure quantities"]
    assert q.max_row == 5 and q.cell(row=q.max_row, column=2).value == "TOTAL"
    assert wb["Summary"]["B3"].value == "Test project" and wb["Sections"].max_row == 6
