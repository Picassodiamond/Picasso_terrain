"""Design data workbook (openpyxl): summary, horizontal, superelevation, vertical, levels,
sections, volumes with live factor formulas, mass haul, section points, structures, checks,
standard parameters."""
from __future__ import annotations

import io
import json
from typing import Any, Sequence

import numpy as np

from ...engine.alignment import format_chainage
from .frame import DrawingInputs

HEAD_FILL = "1F3A5F"
HEAD_FONT = "FFFFFF"


def _style_header(ws, ncols: int, row: int = 1) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color=HEAD_FONT)
        cell.fill = PatternFill("solid", fgColor=HEAD_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _table(wb, name: str, columns: Sequence[str], rows: Sequence[Sequence[Any]], *, formats: dict[int, str] | None = None, start_row: int = 1, ws=None):
    from openpyxl.utils import get_column_letter

    ws = ws or wb.create_sheet(name[:31])
    for j, c in enumerate(columns, 1):
        ws.cell(row=start_row, column=j, value=c)
    _style_header(ws, len(columns), start_row)
    for r in rows:
        ws.append([None if (isinstance(v, float) and not np.isfinite(v)) else v for v in r])
    if formats:
        for j, fmt in formats.items():
            for cell in ws.iter_rows(min_row=start_row + 1, min_col=j, max_col=j):
                cell[0].number_format = fmt
    for j, c in enumerate(columns, 1):
        width = max([len(str(c))] + [len(f"{r[j - 1]:.3f}" if isinstance(r[j - 1], float) else str(r[j - 1])) for r in rows if j - 1 < len(r) and r[j - 1] is not None] or [8])
        ws.column_dimensions[get_column_letter(j)].width = min(max(width + 2, 9), 48)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return ws


def design_workbook(inp: DrawingInputs, *, design: dict, project: dict, sheet_settings: dict | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    al, va = inp.al, inp.va
    ch_fmt, m3, m2, lvl = "0.00", "#,##0.0", "0.000", "0.000"

    # -- summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"{project.get('name', '')} - {design.get('name', '')}"
    ws["A1"].font = Font(bold=True, size=14)
    ctx = inp.context
    kv = [("Project", project.get("name", "")), ("Design", design.get("name", "")), ("CRS", inp.crs_name or project.get("crs", "")),
          ("Standard", f"{inp.standard.get('name', inp.standard.get('id', ''))} ({inp.standard.get('status', 'placeholder')})"),
          ("Road class / terrain / surface", f"{ctx.get('road_class', '')} / {ctx.get('terrain', '')} / {ctx.get('surface', '')}"),
          ("Design speed (km/h)", ctx.get("design_speed")), ("Start chainage", al.start_chainage), ("End chainage", al.end_chainage), ("Length (m)", al.length),
          ("IPs", len(al.ips)), ("Curves", sum(1 for g in al.geometry if g.curve_length > 0)), ("PVIs", len(va.pvis) if va else 0),
          ("Corridor sections", len(inp.sections)), ("Section interval (m)", inp.corridor_params.get("interval")),
          ("Cut (m3)", inp.totals.get("cut")), ("Fill (m3)", inp.totals.get("fill")), ("Adjusted cut (m3)", inp.totals.get("cut_adjusted")),
          ("Adjusted fill (m3)", inp.totals.get("fill_adjusted")), ("Net (m3)", inp.totals.get("net")), ("Structures", len(inp.structures)),
          ("Checks failing", sum(1 for c in inp.checks if c.get("ok") is False)), ("Checks not evaluated", sum(1 for c in inp.checks if c.get("ok") is None))]
    for i, (k, v) in enumerate(kv, 3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    if inp.standard.get("status", "placeholder") != "verified":
        ws.cell(row=len(kv) + 4, column=1, value="Standard values are placeholders: verify against the printed standard before use.").font = Font(italic=True, color="9C4A00")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 44

    # -- horizontal
    rows = []
    for r in inp.curve_rows:
        rows.append([r["label"], r["x"], r["y"], r["chainage"], format_chainage(r["chainage"]), r["radius"] or None, r["transition"] or None, r["deflection_deg"], r.get("turn", ""),
                     r["tangent"] or None, r["arc_length"] or None, r["total_length"] or None, r["external"] or None, r.get("spiral_angle_deg") or None, r.get("shift") or None,
                     r.get("ts"), r.get("sc"), r.get("mc"), r.get("cs"), r.get("st"), r.get("superelevation_pct"), "ok" if r.get("valid", True) else r.get("message", "")])
    _table(wb, "Horizontal", ["IP", "X", "Y", "IP chainage", "Chainage text", "R (m)", "Ls (m)", "Deflection (deg)", "Turn", "T (m)", "Lc (m)", "Total L (m)", "E (m)",
                              "Spiral angle (deg)", "Shift p (m)", "TS/BC", "SC", "MC", "CS", "ST/EC", "e (%)", "Status"], rows,
           formats={2: lvl, 3: lvl, 4: ch_fmt, 8: "0.0000", 10: lvl, 11: lvl, 12: lvl, 13: lvl, 14: "0.0000", 15: "0.0000", 16: ch_fmt, 17: ch_fmt, 18: ch_fmt, 19: ch_fmt, 20: ch_fmt})
    if inp.supere is not None:
        step = max(al.length / 400.0, 5.0)
        se = inp.supere.table(step, al.start_chainage, al.end_chainage)
        _table(wb, "Superelevation", ["Chainage", "Chainage text", "Left slope (%)", "Right slope (%)"],
               [[r["chainage"], format_chainage(r["chainage"]), r["left"], r["right"]] for r in se], formats={1: ch_fmt, 3: "0.00", 4: "0.00"})
        cur = [[c.ip_index, c.radius, c.e_pct, "yes" if c.feasible else "no", c.runoff, c.start, c.full_from, c.full_to, c.end, "left" if c.turn > 0 else "right"] for c in inp.supere.curves]
        _table(wb, "Superelevation curves", ["IP", "R (m)", "e (%)", "Feasible", "Runoff (m)", "Start", "Full from", "Full to", "End", "Turn"], cur,
               formats={5: ch_fmt, 6: ch_fmt, 7: ch_fmt, 8: ch_fmt, 9: ch_fmt})

    # -- vertical
    if va is not None:
        rows = [[r["index"], r["chainage"], format_chainage(r["chainage"]), r["elevation"], r["length"] or None, r.get("grade_in"), r.get("grade_out"), r.get("A"), r.get("K"),
                 (r.get("kind") or ""), r.get("bvc"), r.get("z_bvc"), r.get("evc"), r.get("z_evc"), r.get("turning_chainage"), r.get("turning_elevation"),
                 "ok" if r.get("valid", True) else r.get("message", "")] for r in va.table()]
        _table(wb, "Vertical", ["PVI", "Chainage", "Chainage text", "RL (m)", "L (m)", "Grade in (%)", "Grade out (%)", "A (%)", "K", "Type", "BVC", "RL BVC", "EVC", "RL EVC",
                                "High/low CH", "High/low RL", "Status"], rows,
               formats={2: ch_fmt, 4: lvl, 6: "0.00", 7: "0.00", 8: "0.00", 9: "0.0", 11: ch_fmt, 12: lvl, 13: ch_fmt, 14: lvl, 15: ch_fmt, 16: lvl})

    # -- levels along the centreline
    if inp.sections:
        chs = np.array([s["chainage"] for s in inp.sections], float)
    else:
        chs = al.stations(20.0, include_curve_points=True)
    gz = inp.ground_z(chs)
    dz = va.elevation_at(np.clip(chs, va.start_chainage, va.end_chainage)) if va is not None else np.full(len(chs), np.nan)
    gr = va.grade_at(np.clip(chs, va.start_chainage, va.end_chainage)) if va is not None else np.full(len(chs), np.nan)
    rows = []
    for i, c in enumerate(chs):
        x, y = al.point_at(float(c))
        sl = inp.supere.slopes_at(float(c)) if inp.supere is not None else (None, None)
        rows.append([float(c), format_chainage(float(c)), x, y, float(gz[i]) if np.isfinite(gz[i]) else None, float(dz[i]) if np.isfinite(dz[i]) else None,
                     float(dz[i] - gz[i]) if np.isfinite(dz[i]) and np.isfinite(gz[i]) else None, float(gr[i]) if np.isfinite(gr[i]) else None, sl[0], sl[1]])
    _table(wb, "Levels", ["Chainage", "Chainage text", "X", "Y", "Ground RL", "Design RL", "Fill (+) / Cut (-)", "Grade (%)", "Left slope (%)", "Right slope (%)"], rows,
           formats={1: ch_fmt, 3: lvl, 4: lvl, 5: lvl, 6: lvl, 7: lvl, 8: "0.00", 9: "0.00", 10: "0.00"})

    # -- sections, volumes, mass haul, section points
    if inp.sections:
        secs = sorted(inp.sections, key=lambda s: s["chainage"])
        rows = [[s["chainage"], format_chainage(s["chainage"]), s.get("template_id"), s.get("design_z"), s.get("ground_z"), s["cut_area"], s["fill_area"],
                 s["left"].get("kind"), s["left"].get("hinge_offset"), s["left"].get("height"), s["left"].get("catch_offset"), s["left"].get("benches"),
                 s["right"].get("kind"), s["right"].get("hinge_offset"), s["right"].get("height"), s["right"].get("catch_offset"), s["right"].get("benches"),
                 ", ".join(s.get("flags") or [])] for s in secs]
        _table(wb, "Sections", ["Chainage", "Chainage text", "Template", "Design RL", "Ground RL", "Cut area (m2)", "Fill area (m2)", "Left", "L hinge offset", "L height", "L catch offset",
                                "L benches", "Right", "R hinge offset", "R height", "R catch offset", "R benches", "Flags"], rows,
               formats={1: ch_fmt, 4: lvl, 5: lvl, 6: m2, 7: m2, 9: "0.00", 10: "0.00", 11: "0.00", 14: "0.00", 15: "0.00", 16: "0.00"})

        ws = wb.create_sheet("Volumes")
        cf = float(inp.corridor_params.get("cut_factor") or 1.0)
        ff = float(inp.corridor_params.get("fill_factor") or 1.0)
        ws["A1"], ws["B1"] = "Cut factor (shrinkage of cut placed as fill)", cf
        ws["A2"], ws["B2"] = "Fill factor (compaction demand)", ff
        ws["A1"].font = ws["A2"].font = Font(bold=True)
        ws["D1"] = "Edit the factors: adjusted volumes, net and cumulative recalculate."
        ws["D1"].font = Font(italic=True)
        by_ch = {round(s["chainage"], 3): s for s in secs}
        cols = ["From", "To", "Length (m)", "Cut area from", "Cut area to", "Fill area from", "Fill area to", "Cut (m3)", "Fill (m3)", "Adj. cut (m3)", "Adj. fill (m3)", "Net (m3)", "Cumulative (m3)", "Method"]
        start = 4
        for j, c in enumerate(cols, 1):
            ws.cell(row=start, column=j, value=c)
        _style_header(ws, len(cols), start)
        r = start + 1
        for v in inp.volumes:
            a, b = by_ch.get(round(v["from"], 3), {}), by_ch.get(round(v["to"], 3), {})
            ws.append([v["from"], v["to"], v["length"], a.get("cut_area"), b.get("cut_area"), a.get("fill_area"), b.get("fill_area"), v["cut"], v["fill"],
                       f"=H{r}*$B$1", f"=I{r}*$B$2", f"=J{r}-K{r}", f"=L{r}" if r == start + 1 else f"=M{r - 1}+L{r}", v.get("method", "")])
            r += 1
        last = r - 1
        if last >= start + 1:
            ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
            for col in "HIJKL":
                ws[f"{col}{r}"] = f"=SUM({col}{start + 1}:{col}{last})"
                ws[f"{col}{r}"].font = Font(bold=True)
        for col, fmt in (("A", ch_fmt), ("B", ch_fmt), ("C", "0.00"), ("D", m2), ("E", m2), ("F", m2), ("G", m2), ("H", m3), ("I", m3), ("J", m3), ("K", m3), ("L", m3), ("M", m3)):
            for cell in ws[col][start:]:
                cell.number_format = fmt
            ws.column_dimensions[col].width = 14
        ws.column_dimensions["A"].width = 16
        ws.freeze_panes = ws.cell(row=start + 1, column=1)

        _table(wb, "Mass haul", ["Chainage", "Chainage text", "Cumulative (m3)"], [[m["chainage"], format_chainage(m["chainage"]), m["cumulative"]] for m in inp.mass_haul],
               formats={1: ch_fmt, 3: m3})

        rows = []
        for s in secs:
            g = np.asarray(s.get("ground") or [], float)
            for o, z in s.get("design") or []:
                zg = float(np.interp(o, g[:, 0], g[:, 1])) if len(g) and g[:, 0].min() - 1e-6 <= o <= g[:, 0].max() + 1e-6 else None
                rows.append([s["chainage"], format_chainage(s["chainage"]), "design", o, z, zg])
            for o, z in s.get("ground") or []:
                rows.append([s["chainage"], format_chainage(s["chainage"]), "ground", o, None, z])
        _table(wb, "Section points", ["Chainage", "Chainage text", "Line", "Offset (m)", "Design RL", "Ground RL"], rows, formats={1: ch_fmt, 4: "0.000", 5: lvl, 6: lvl})

    # -- structures and their quantities
    from ...design.structures import MATERIALS, STRUCTURE_KINDS, bill_of_quantities, type_spec

    rows = []
    for s in inp.structures:
        p = s.get("params") or {}
        spec = STRUCTURE_KINDS.get(s["kind"], {})
        t = str(p.get("type") or "")
        rows.append([s["id"], spec.get("label", s["kind"]), type_spec(s["kind"], t).get("label", t), s.get("side") or "", s["from"], s["to"], s["to"] - s["from"],
                     format_chainage(s["from"]), format_chainage(s["to"]), p.get("height"), p.get("span"), p.get("width"), p.get("depth"),
                     s.get("source", ""), json.dumps({k: v for k, v in p.items() if k != "type"}), s.get("note", "")])
    _table(wb, "Structures", ["Id", "Kind", "Type", "Side", "From", "To", "Length (m)", "From text", "To text", "Height (m)", "Span (m)", "Width (m)", "Depth (m)",
                              "Source", "Parameters", "Note"], rows, formats={5: ch_fmt, 6: ch_fmt, 7: "0.00", 10: "0.00", 11: "0.00", 12: "0.00", 13: "0.00"})
    if inp.structures:
        boq = bill_of_quantities(inp.structures)
        mats = sorted(boq["totals"])
        rows = []
        for r in boq["rows"]:
            kind = STRUCTURE_KINDS.get(r["kind"], {}).get("label", r["kind"])
            rows.append([r["id"], kind, type_spec(r["kind"], str(r["type"] or "")).get("label", r["type"]), r["length"], r["sides"], r["count"]]
                        + [r["total"].get(m) for m in mats])
        rows.append(["", "TOTAL", "", None, None, None] + [boq["totals"].get(m) for m in mats])
        cols = ["Id", "Kind", "Type", "Length (m)", "Sides", "Count"] + [f"{MATERIALS.get(m, {}).get('label', m)} ({MATERIALS.get(m, {}).get('unit', '')})" for m in mats]
        _table(wb, "Structure quantities", cols, rows, formats={i: "#,##0.00" for i in range(7, 7 + len(mats))})

    # -- checks and standard
    rows = [[c.get("stage", ""), c.get("id"), c.get("label"), c.get("where"), c.get("actual"), c.get("limit"), c.get("unit"),
             "ok" if c.get("ok") else ("FAIL" if c.get("ok") is False else "not evaluated"), c.get("severity"), c.get("status"), c.get("source"), c.get("message")] for c in inp.checks]
    _table(wb, "Checks", ["Stage", "Check", "Label", "Where", "Actual", "Limit", "Unit", "Result", "Severity", "Parameter status", "Source", "Message"], rows, formats={5: "0.000", 6: "0.000"})
    rows = [[p.get("key"), p.get("value"), p.get("unit"), p.get("status"), p.get("source"), p.get("note")] for p in inp.standard_params]
    _table(wb, "Standard", ["Parameter", "Value", "Unit", "Status", "Source", "Note"], rows)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
